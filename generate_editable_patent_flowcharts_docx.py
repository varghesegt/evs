import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

FLOWCHARTS = [
    {
        "fig_num": "FIG. 1",
        "sheet": "SHEET 1 OF 6",
        "title": "DISTRIBUTED CYBER-PHYSICAL SYSTEM ARCHITECTURE FLOW",
        "field": "Technical Field: Distributed multi-tier cloud computing architecture integrating non-contact volumetric vision, dynamic perishability kinetics, combinatorial logistics, and cryptographic environmental ledgers.",
        "steps": [
            "1. Problem Definition & Objective Setting",
            "2. Non-Contact Image Capture & GPS Geocoding",
            "3. Calibrated Container Geometry & Fluid Volume Calculation",
            "4. 12-Category Food Bulk Density Physics Transformation",
            "5. Dynamic Shelf-Life Decay & Microbiological Urgency Modeling",
            "6. Spatial Clustering & Proactive Shelter Deficit Forecasting",
            "7. Perishability-Constrained Capacitated Vehicle Routing (CVRP)",
            "8. Physical Handover via Single-Use Cryptographic QR Token",
            "9. Zero-Trust Verification & Secure State Transition",
            "10. IPCC AR6 Multi-Gas Impact & Methane GWP 28× Calculation",
            "11. Immutable SHA-256 Ledger Commitment & Certificate Issuance"
        ],
        "summary": "FIG. 1 depicts the end-to-end cyber-physical architecture flow, orchestrating client image capture, container volume calibration, food density physics, microbiological shelf-life modeling, CVRP vehicle routing, and cryptographic SHA-256 environmental ledger accounting into a unified automated redistribution pipeline."
    },
    {
        "fig_num": "FIG. 2",
        "sheet": "SHEET 2 OF 6",
        "title": "NON-CONTACT VOLUMETRIC FOOD MASS ESTIMATION & ERROR-BOUNDING",
        "field": "Technical Field: Computer vision, container geometry modeling, bulk density physics tensors, and closed-loop statistical uncertainty bounding.",
        "steps": [
            "1. Ingestion of Multi-View Images & Container Metadata",
            "2. Calibrated Receptacle Reference Volume Matrix Lookup",
            "3. Surface Fill Percentage Detection & Fluid Volume Calculation",
            "4. 12-Category Bulk Density Tensor Mapping (kg/Liter)",
            "5. Midpoint Physical Mass & Portion Count Determination",
            "6. Perspective Assessment (Orthogonal Dual-View vs. Monocular)",
            "7. Dynamic Statistical Uncertainty Bounding (±8% vs. ±12%)",
            "8. Standard Receptacle Repeatability Confidence Boost (+5%)",
            "9. Synthesis of Bounded Mass Tuple [Min Mass, Max Mass]",
            "10. Real-World Verified Physical Weight Check-in at Shelter Scale",
            "11. Closed-Loop Continuous MAPE Model Recalibration"
        ],
        "summary": "FIG. 2 illustrates the volumetric mass estimation method. Grounding visual fill percentages in standardized catering receptacle geometry and category bulk density eliminates monocular scale ambiguity. Dual-perspective capture bounds uncertainty to ±8%, while closed-loop MAPE feedback dynamically recalibrates historical model accuracy."
    },
    {
        "fig_num": "FIG. 3",
        "sheet": "SHEET 3 OF 6",
        "title": "TRI-FACTOR MICROBIOLOGICAL DECAY & PRIORITY URGENCY SCORING",
        "field": "Technical Field: Food microbiology kinetics, room-temperature shelf-life curves, and multi-objective priority optimization.",
        "steps": [
            "1. Donation Preparation Timestamp & Category Registration",
            "2. Category-Specific Room-Temperature Shelf-Life Matrix Query",
            "3. Remaining Safe Distribution Time Calculation (t_rem)",
            "4. Mass Impact Normalization Vector Computation (S_qty)",
            "5. Biochemical Freshness Urgency Vector Computation (S_fresh)",
            "6. Elapsed Preparation Time Decay Vector Computation (S_time)",
            "7. Tri-Factor Weighted Composite Priority Optimization",
            "8. Spoilage Threshold Verification (t_rem > 0)",
            "9. Emergency Dispatch Alert Broadcast to Ranked Nearby NGOs",
            "10. Bio-Hazard Diversion to Composting for Expired Batches",
            "11. Persistence of Priority Urgency Record in Database Registry"
        ],
        "summary": "FIG. 3 details the microbiological decay logic. The priority engine couples category bacterial growth kinetics with batch mass and chronological transit deadlines, preventing the redistribution of spoiled food while guaranteeing that high-volume perishable batches are matched to nearby NGOs within minimum latency windows."
    },
    {
        "fig_num": "FIG. 4",
        "sheet": "SHEET 4 OF 6",
        "title": "PERISHABILITY-CONSTRAINED MULTI-STOP CVRP LOGISTICS ROUTING",
        "field": "Technical Field: Operations research, Capacitated Vehicle Routing Problem (CVRP), payload capacity boundaries, and transport emission minimization.",
        "steps": [
            "1. Multi-Stop Dispatch Request & Volunteer Location Ingestion",
            "2. Vehicle Gross Payload Capacity Limit Definition (C_max)",
            "3. Candidate Donor Pickup Nodes & Recipient Shelter Set Assembly",
            "4. All-Pairs Haversine Geospatial Distance Matrix Computation",
            "5. Trajectory Initialization & Cumulative Payload Tracking",
            "6. Spatial Nearest-Neighbor Candidate Node Search",
            "7. Vehicle Payload Capacity Feasibility Assessment (L + M <= C_max)",
            "8. Mandatory Intermediate Delivery Dropoff upon Capacity Breach",
            "9. Topological Precedence Delivery Dropoff Execution",
            "10. Consolidated Route Optimization vs. Radial Baseline Calculation",
            "11. Fuel Savings & Transport CO2e Mitigation Modeling",
            "12. Waypoint Manifest Transmission to Mobile Navigation Client"
        ],
        "summary": "FIG. 4 illustrates the CVRP routing solver. The algorithm resolves multi-stop vehicle trajectories subject to gross payload weight limits and strict chronological pickup-before-delivery dependencies, consolidating disparate donor batches into fuel-minimized routes and quantifying transport emission reductions."
    },
    {
        "fig_num": "FIG. 5",
        "sheet": "SHEET 5 OF 6",
        "title": "PROACTIVE SHELTER DEFICIT & TIME-SERIES DEMAND FORECASTING",
        "field": "Technical Field: Time-series statistical forecasting, recipient inventory telemetry, and proactive commercial surplus solicitation.",
        "steps": [
            "1. 24-Hour Regional Forecast Cycle Initialization",
            "2. Recipient Shelter Inventory & Bed Capacity Telemetry Ingestion",
            "3. Day-of-Week Trend & Local Festival Calendar Multipliers",
            "4. Time-Series Meal Deficit Calculation (Cap × 2.20 × Multipliers)",
            "5. Urban Corridor Geospatial Clustering & High-Demand Zone Demarcation",
            "6. Deficit Urgency Classification & Peak Hunger Window Demarcation",
            "7. Commercial Banquet & Institutional Kitchen Surplus Aggregation",
            "8. Deficit Coverage Percentage Evaluation (Surplus / Deficit × 100)",
            "9. Automated Proactive Dispatch Solicitations (6-Hour Notice)",
            "10. Forward-Looking Surplus Matching to Deficit Corridors",
            "11. Real-Time Demand Matrix Rendering on Administrative Dashboard"
        ],
        "summary": "FIG. 5 demonstrates the proactive demand forecasting architecture. By converting historical day-of-week consumption patterns and shelter storage capacities into forward-looking deficit horizons, the system solicits commercial surplus hours before waste occurs, mitigating supply shortages in critical urban feeding corridors."
    },
    {
        "fig_num": "FIG. 6",
        "sheet": "SHEET 6 OF 6",
        "title": "ZERO-TRUST CHAIN OF CUSTODY & CRYPTOGRAPHIC SHA-256 AUDIT LEDGER",
        "field": "Technical Field: Cryptographic authentication, zero-trust token handshakes, IPCC AR6 multi-gas environmental accounting, and immutable SHA-256 ledgers.",
        "steps": [
            "1. Validated Food Donation Pipeline Assignment",
            "2. Non-Deterministic Single-Use Handover Token Generation (τ)",
            "3. High-Density 2D Optical QR Code Rendering & Container Affixation",
            "4. Volunteer In-Transit Tracking to Destination Shelter Depot",
            "5. Optical Barcode Scan & Token Data Ingestion at Physical Handover",
            "6. Zero-Trust Cryptographic String Comparison (τ_input == τ)",
            "7. Security State Lockout upon Handover Token Mismatch",
            "8. Authorized State Transition to Delivered upon Token Verification",
            "9. IPCC AR6 Multi-Gas Environmental Nexus Accounting",
            "10. Methane Landfill Mitigation (GWP 28×) & Virtual Water Savings",
            "11. Canonical Ledger Payload Serialization: \"EVS-LEDGER-V1|...\"",
            "12. 256-Bit SHA-256 Cryptographic Hash Digest Generation",
            "13. Immutable Ledger Record Persistence & Public Audit API Exposure"
        ],
        "summary": "FIG. 6 depicts the cryptographic verification subsystem. Requiring physical scanning of a single-use non-deterministic token establishes a zero-trust custody gate, eliminating phantom deliveries. Committing the multi-gas environmental nexus to a SHA-256 hash digest guarantees data immutability, preventing carbon offset double-counting."
    }
]

def make_round_rect_box_xml(text, shape_id, cx=5400000, cy=410000, font_size=21):
    escaped_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    xml = f"""
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:pPr>
    <w:jc w:val="center"/>
    <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
  </w:pPr>
  <w:r>
    <w:drawing>
      <wp:inline distT="0" distB="0" distL="0" distR="0">
        <wp:extent cx="{cx}" cy="{cy}"/>
        <wp:docPr id="{shape_id}" name="FlowBox_{shape_id}"/>
        <wps:wsp>
          <wps:cNvSpPr/>
          <wps:spPr>
            <a:xfrm>
              <a:off x="0" y="0"/>
              <a:ext cx="{cx}" cy="{cy}"/>
            </a:xfrm>
            <a:prstGeom prst="roundRect">
              <a:avLst>
                <a:gd name="adj" fmla="val 13500"/>
              </a:avLst>
            </a:prstGeom>
            <a:solidFill>
              <a:srgbClr val="FFFFFF"/>
            </a:solidFill>
            <a:ln w="22000">
              <a:solidFill>
                <a:srgbClr val="000000"/>
              </a:solidFill>
            </a:ln>
          </wps:spPr>
          <wps:txbx>
            <w:txbxContent>
              <w:p>
                <w:pPr>
                  <w:jc w:val="center"/>
                  <w:spacing w:before="0" w:after="0"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Cambria" w:hAnsi="Cambria" w:cs="Cambria"/>
                    <w:b/>
                    <w:sz w:val="{font_size}"/>
                    <w:color w:val="000000"/>
                  </w:rPr>
                  <w:t>{escaped_text}</w:t>
                </w:r>
              </w:p>
            </w:txbxContent>
          </wps:txbx>
          <wps:bodyPr vert="horz" lIns="72000" tIns="36000" rIns="72000" bIns="36000" anchor="ctr"/>
        </wps:wsp>
      </wp:inline>
    </w:drawing>
  </w:r>
</w:p>
"""
    return parse_xml(xml)

def make_down_arrow_xml(arrow_id, cx=190000, cy=170000):
    xml = f"""
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:pPr>
    <w:jc w:val="center"/>
    <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
  </w:pPr>
  <w:r>
    <w:drawing>
      <wp:inline distT="0" distB="0" distL="0" distR="0">
        <wp:extent cx="{cx}" cy="{cy}"/>
        <wp:docPr id="{arrow_id}" name="Arrow_{arrow_id}"/>
        <wps:wsp>
          <wps:cNvSpPr/>
          <wps:spPr>
            <a:xfrm>
              <a:off x="0" y="0"/>
              <a:ext cx="{cx}" cy="{cy}"/>
            </a:xfrm>
            <a:prstGeom prst="downArrow">
              <a:avLst>
                <a:gd name="adj1" fmla="val 50000"/>
                <a:gd name="adj2" fmla="val 50000"/>
              </a:avLst>
            </a:prstGeom>
            <a:solidFill>
              <a:srgbClr val="000000"/>
            </a:solidFill>
            <a:ln w="9525">
              <a:solidFill>
                <a:srgbClr val="000000"/>
              </a:solidFill>
            </a:ln>
          </wps:spPr>
          <wps:bodyPr/>
        </wps:wsp>
      </wp:inline>
    </w:drawing>
  </w:r>
</w:p>
"""
    return parse_xml(xml)

def generate_editable_document():
    doc = Document()
    
    # 0.55" top/bottom margins, 0.70" left/right margins for full page utilization
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.70)
        section.right_margin = Inches(0.70)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("PATENT APPLICATION SPECIFICATION DRAWINGS • EDITABLE WORD FORMAT")
        hr.font.name = "Cambria"
        hr.font.size = Pt(8.5)
        hr.font.bold = True
        hr.font.color.rgb = RGBColor(0, 0, 0)

    shape_counter = 1000

    for idx, item in enumerate(FLOWCHARTS):
        num_steps = len(item["steps"])
        
        # Adaptive dimensions based on step count so each figure perfectly fits 1 page
        if num_steps >= 13:
            box_cy = 350000  # ~0.38 in
            arrow_cy = 135000 # ~0.15 in
            font_sz = 19     # 9.5 pt
        elif num_steps >= 12:
            box_cy = 380000  # ~0.41 in
            arrow_cy = 150000 # ~0.16 in
            font_sz = 20     # 10 pt
        else:
            box_cy = 410000  # ~0.45 in
            arrow_cy = 165000 # ~0.18 in
            font_sz = 21     # 10.5 pt

        # 1. Sheet Number (Right-aligned)
        p_sheet = doc.add_paragraph()
        p_sheet.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sheet.paragraph_format.space_before = Pt(0)
        p_sheet.paragraph_format.space_after = Pt(1)
        r_sheet = p_sheet.add_run(item["sheet"])
        r_sheet.font.name = "Cambria"
        r_sheet.font.size = Pt(9.0)
        r_sheet.font.bold = True
        r_sheet.font.color.rgb = RGBColor(0, 0, 0)
        
        # 2. Figure Title (Centered)
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(1)
        r_title = p_title.add_run(f"{item['fig_num']}: {item['title']}")
        r_title.font.name = "Cambria"
        r_title.font.size = Pt(12.0)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0, 0, 0)
        
        # 3. Technical Field (Centered, Italic)
        p_field = doc.add_paragraph()
        p_field.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_field.paragraph_format.space_before = Pt(0)
        p_field.paragraph_format.space_after = Pt(3)
        r_field = p_field.add_run(item["field"])
        r_field.font.name = "Cambria"
        r_field.font.size = Pt(8.0)
        r_field.font.italic = True
        r_field.font.color.rgb = RGBColor(0, 0, 0)
        
        # 4. Add Editable Rounded-Box Flowchart Steps & Arrows
        for s_idx, step_text in enumerate(item["steps"]):
            shape_counter += 1
            box_xml = make_round_rect_box_xml(step_text, shape_counter, cx=5400000, cy=box_cy, font_size=font_sz)
            doc._body._body.append(box_xml)
            
            if s_idx < num_steps - 1:
                shape_counter += 1
                arrow_xml = make_down_arrow_xml(shape_counter, cx=190000, cy=arrow_cy)
                doc._body._body.append(arrow_xml)
        
        # 5. Technical Description Paragraph
        p_desc = doc.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_desc.paragraph_format.space_before = Pt(4)
        p_desc.paragraph_format.space_after = Pt(0)
        p_desc.paragraph_format.line_spacing = 1.05
        
        r_lbl = p_desc.add_run(f"Technical Description of {item['fig_num']}: ")
        r_lbl.font.name = "Cambria"
        r_lbl.font.size = Pt(8.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(0, 0, 0)
        
        r_desc = p_desc.add_run(item["summary"])
        r_desc.font.name = "Cambria"
        r_desc.font.size = Pt(8.5)
        r_desc.font.color.rgb = RGBColor(0, 0, 0)
        
        # Page Break between figures (Pages 1 to 5)
        if idx < len(FLOWCHARTS) - 1:
            doc.add_page_break()

    output_filename = "EcoShare_AI_Patent_Architecture_Flowcharts.docx"
    doc.save(output_filename)
    print(f"Successfully generated editable Word document: {output_filename}")

if __name__ == "__main__":
    generate_editable_document()
